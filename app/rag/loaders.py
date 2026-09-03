"""Document loaders for ingesting marketing knowledge into the RAG pipeline.

Each loader returns a `LoadedDocument` with raw text content and metadata,
ready for chunking + embedding. `source_type` matches the payload index
already set up in Qdrant (app/qdrant/service.py).
"""
from __future__ import annotations

import csv
import uuid
from dataclasses import dataclass, field
from pathlib import Path

import httpx
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader


@dataclass
class LoadedDocument:
    """A single loaded document ready for chunking + embedding."""

    content: str
    source_type: str  # "pdf" | "docx" | "txt" | "csv" | "web"
    document_id: str = field(default_factory=lambda: str(uuid.uuid4()))
    metadata: dict = field(default_factory=dict)


class DocumentLoadError(Exception):
    """Raised when a document cannot be read or parsed."""


def load_pdf(path: str | Path) -> LoadedDocument:
    path = Path(path)
    try:
        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise DocumentLoadError(f"Failed to read PDF {path}: {exc}") from exc

    if not text.strip():
        raise DocumentLoadError(f"No extractable text found in PDF {path}")

    return LoadedDocument(
        content=text,
        source_type="pdf",
        metadata={"filename": path.name, "pages": len(reader.pages)},
    )


def load_docx(path: str | Path) -> LoadedDocument:
    path = Path(path)
    try:
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as exc:
        raise DocumentLoadError(f"Failed to read DOCX {path}: {exc}") from exc

    if not text.strip():
        raise DocumentLoadError(f"No extractable text found in DOCX {path}")

    return LoadedDocument(
        content=text,
        source_type="docx",
        metadata={"filename": path.name},
    )


def load_txt(path: str | Path, encoding: str = "utf-8") -> LoadedDocument:
    path = Path(path)
    try:
        text = path.read_text(encoding=encoding)
    except Exception as exc:
        raise DocumentLoadError(f"Failed to read TXT {path}: {exc}") from exc

    if not text.strip():
        raise DocumentLoadError(f"Empty text file {path}")

    return LoadedDocument(content=text, source_type="txt", metadata={"filename": path.name})


def load_csv(path: str | Path, text_columns: list[str] | None = None) -> LoadedDocument:
    """Flatten a CSV into a single text blob, one line per row."""
    path = Path(path)
    rows_text: list[str] = []
    try:
        with path.open(newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                cols = text_columns or list(row.keys())
                line = " | ".join(str(row.get(c, "")) for c in cols)
                if line.strip(" |"):
                    rows_text.append(line)
    except Exception as exc:
        raise DocumentLoadError(f"Failed to read CSV {path}: {exc}") from exc

    text = "\n".join(rows_text)
    if not text.strip():
        raise DocumentLoadError(f"No rows found in CSV {path}")

    return LoadedDocument(
        content=text,
        source_type="csv",
        metadata={"filename": path.name, "row_count": len(rows_text)},
    )


async def load_web(url: str, timeout: float = 15.0) -> LoadedDocument:
    """Fetch a URL and extract its visible text content."""
    try:
        async with httpx.AsyncClient(timeout=timeout) as client:
            response = await client.get(url, follow_redirects=True)
            response.raise_for_status()
    except Exception as exc:
        raise DocumentLoadError(f"Failed to fetch URL {url}: {exc}") from exc

    soup = BeautifulSoup(response.text, "lxml")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()

    text = " ".join(soup.get_text(separator=" ").split())
    if not text.strip():
        raise DocumentLoadError(f"No extractable text found at {url}")

    return LoadedDocument(content=text, source_type="web", metadata={"url": url})


LOADER_BY_EXTENSION = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".txt": load_txt,
    ".csv": load_csv,
}


def load_file(path: str | Path) -> LoadedDocument:
    """Dispatch to the correct loader based on file extension."""
    path = Path(path)
    loader = LOADER_BY_EXTENSION.get(path.suffix.lower())
    if loader is None:
        raise DocumentLoadError(f"Unsupported file type: {path.suffix}")
    return loader(path)